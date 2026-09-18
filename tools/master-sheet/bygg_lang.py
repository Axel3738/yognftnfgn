"""Alla batcher i EN lång flik, under varandra (Axels beslut 2026-09-17: scrolla + Ctrl+F).

Varje batch behåller sina egna rubrikrader, sin formatering, sina bilder och
länkar. Före varje batch ligger en blå bannerrad med batchens namn. Inga
kolumner döljs (Batch 8 och Kalenderkungen hade dolda kolumner i originalet).
"""
import io
import re
import subprocess
from copy import copy, deepcopy

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter, column_index_from_string
from PIL import Image as PILImage

import datetime
from bygg_master import KALLOR, krymp, hamta_url, IMAGE_RE

# Batch 6 kom tillbaka ifylld från leverantören 2026-09-18 (taköverdraget: nio storlekar)
KALLOR = [("06-ny.xlsx" if fil == "06.xlsx" else fil, namn) for fil, namn in KALLOR]

NYA_MARKNADER = [("CANADA", "FFC9DAF8"), ("AUSTRALIA", "FFFFE599"), ("NEW ZEALAND", "FFB6D7A8")]
UNDERRUBRIK = ["Qty", "Product cost", "Shipping cost", "Total ex. tax", "Delivery time", "Shipping method"]
tidsfel = 0


def tid_till_tal(v):
    """Google Sheets tolkade leverantörens "17.42" som klockslaget 17:42 i US-kolumnerna.
    Tillbaka: timmar.minuter (2 days, 4:27 → 52.27). Kontrollerat mot SE-kolumnen: J19 52.27 = AQ19 52:27."""
    if isinstance(v, datetime.time):
        return float(f"{v.hour}.{v.minute:02d}")
    if isinstance(v, datetime.timedelta):
        h, rest = divmod(int(v.total_seconds()), 3600)
        return float(f"{h}.{rest // 60:02d}")
    if isinstance(v, datetime.datetime):
        return float(f"{v.hour}.{v.minute:02d}")
    return v


def _num(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _kand(p):
    """Ett tal som lästs som klockslag kan ha varit p eller p−0.40 (28.79 → 28:79 → 29:19)."""
    return [round(p, 2), round(p - 0.40, 2)] if isinstance(p, float) else []


def ratta_tidsvarden(srcv):
    """Hittar marknadsblock (rad 3 rubrik, rad 4 Qty/Product cost/Shipping cost/Total ex. tax)
    där leverantörens decimaltal lästs som klockslag, och räknar tillbaka dem med SE-blockets
    Product cost som facit (produktkostnaden är samma i alla marknader) och total = produkt + frakt.
    Returnerar {(rad, kol): (värde, kommentar-eller-None)}."""
    ut = {}
    under = {c.column: str(c.value).strip().lower() for c in srcv[4] if c.value}
    rub3 = {c.column: str(c.value).strip().lower() for c in srcv[3] if c.value}
    se_kost = next((k for k, v in rub3.items() if v == "product cost"), None)
    block = []
    for k, v in under.items():
        if v == "qty":
            block.append(dict(kost=k + 1, frakt=k + 2, tot=k + 3))
    for r in range(5, srcv.max_row + 1):
        for b in block:
            celler = {n: srcv.cell(r, k).value for n, k in b.items()}
            if not any(isinstance(x, (datetime.time, datetime.timedelta, datetime.datetime)) for x in celler.values()):
                continue
            C = _num(srcv.cell(r, se_kost).value) if se_kost else None
            pk, pf, pt = (tid_till_tal(celler[n]) if isinstance(celler[n], (datetime.time, datetime.timedelta, datetime.datetime)) else _num(celler[n]) for n in ("kost", "frakt", "tot"))
            # produktkostnad
            if C is not None and C in _kand(pk):
                ut[(r, b["kost"])] = (C, None)
            elif pk is not None:
                ut[(r, b["kost"])] = (pk, f"Read as a time by Google Sheets; could be {pk:.2f} or {pk - 0.40:.2f}. SE product cost is {C}.")
            # frakt + total
            losn = [(f, t) for f in _kand(pf) for t in _kand(pt) if C is not None and abs(C + f - t) < 0.006]
            if len(losn) == 1:
                ut[(r, b["frakt"])] = (losn[0][0], None)
                ut[(r, b["tot"])] = (losn[0][1], None)
            else:
                alt = " or ".join(f"{f:.2f} / {t:.2f}" for f, t in losn) if losn else "no combination matches product + shipping = total"
                note = f"UNCERTAIN — read as a time by Google Sheets. Shipping / total could be {alt}. Please re-enter as plain numbers."
                if pf is not None:
                    ut[(r, b["frakt"])] = (losn[0][0] if losn else pf, note)
                if pt is not None:
                    ut[(r, b["tot"])] = (losn[0][1] if losn else pt, note)
    return ut


def lagg_till_marknader(ws, src, offset, sista, rubrikrad, underrad):
    """Tre nya ifyllnadsblock (CA, AU, NZ) direkt efter batchens sista marknadsblock,
    med samma stil, samma kolumnbredder och Qty 1/2/3 förifyllt som i UK-blocket."""
    rub = [c for c in src[rubrikrad] if c.value not in (None, "")]
    if not rub:
        return None
    sista_kol = max(c.column for c in rub)
    # blocken är 6 kolumner breda; UK-blocket är förlagan
    uk = next((c.column for c in rub if str(c.value).strip().upper() == "UK"), None)
    if uk is None:
        return None
    qty_kolumner = [c.column for c in src[underrad] if str(c.value).strip().lower() == "qty"] if underrad else [c.column for c in rub if c.column >= 17]
    # börja efter sista blocket OCH efter varje sammanslagen cell som sticker ut åt höger
    start = max(sista_kol + 6, max((m.max_col for m in src.merged_cells.ranges), default=0) + 1)
    for i, (namn, farg) in enumerate(NYA_MARKNADER):
        k0 = start + i * 6
        for j in range(6):
            kol = k0 + j
            ws.column_dimensions[get_column_letter(kol)].width = src.column_dimensions[get_column_letter(uk + j)].width or 12
            for r in range(rubrikrad, sista + 1):
                forlaga = src.cell(r, uk + j)
                d = ws.cell(r + offset, kol)
                if type(d).__name__ == "MergedCell":
                    continue
                if forlaga.has_style:
                    d.font, d.fill, d.border, d.alignment, d.number_format = (
                        copy(forlaga.font), copy(forlaga.fill), copy(forlaga.border), copy(forlaga.alignment), forlaga.number_format)
                if r == rubrikrad:
                    if j == 0:
                        d.value = namn
                    d.fill = PatternFill("solid", fgColor=farg)
                elif underrad and r == underrad:
                    d.value = UNDERRUBRIK[j]
                elif j == 0:
                    q = next((src.cell(r, q).value for q in qty_kolumner if src.cell(r, q).value not in (None, "")), None)
                    d.value = q
        ws.merge_cells(start_row=rubrikrad + offset, end_row=rubrikrad + offset, start_column=k0, end_column=k0 + 5)
    return start

UT = "Master sheet - alla produkter (en flik).xlsx"
BANNER_FILL = PatternFill("solid", fgColor="1F4E78")
BANNER_FONT = Font(name="Arial", bold=True, size=14, color="FFFFFF")

master = openpyxl.Workbook()
ws = master.active
ws.title = "ALLA PRODUKTER"
bredd = {}
rad = 1
innehall = []  # (bannerrad, namn)
logg = []

for fil, namn in KALLOR:
    src = openpyxl.load_workbook(fil).active
    srcv = openpyxl.load_workbook(fil, data_only=True).active

    # bannerrad
    b = ws.cell(rad, 1, f"▶ {namn.upper()}")
    b.font, b.fill = BANNER_FONT, BANNER_FILL
    for c in range(2, 80):
        ws.cell(rad, c).fill = BANNER_FILL
    ws.row_dimensions[rad].height = 24
    innehall.append((rad, namn))
    offset = rad  # källans rad 1 hamnar på rad offset+1

    # Google exporterar 1 000 rader per ark — klipp vid sista raden med innehåll
    sista = max(
        [r[0].row for r in src.iter_rows() if any(c.value not in (None, "") for c in r)]
        + [rng.max_row for rng in src.merged_cells.ranges]
        + [img.anchor._from.row + 3 for img in src._images]
    )
    n_img = 0
    ratt = ratta_tidsvarden(srcv) if namn not in ("Batch 1", "Batch 2") else {}
    osakra = 0
    for row in src.iter_rows(max_row=sista):
        for c in row:
            v = srcv[c.coordinate].value
            rattad = isinstance(v, (datetime.time, datetime.timedelta, datetime.datetime))
            kommentar = None
            if (c.row, c.column) in ratt:
                v, kommentar = ratt[(c.row, c.column)]
                rattad = True; tidsfel += 1
            elif rattad:
                v = tid_till_tal(v); tidsfel += 1
            raw = c.value
            if isinstance(raw, str) and raw.startswith("="):
                m = IMAGE_RE.match(raw)
                if m:
                    pil = hamta_url(m.group(1)) or hamta_url(m.group(1))  # ett omförsök
                    if pil is not None:
                        xi = XlImage(krymp(pil))
                        xi.width, xi.height = 120, 90
                        xi.anchor = f"{get_column_letter(c.column)}{c.row + offset}"
                        ws.add_image(xi)
                        n_img += 1
                        v = None
                    else:
                        v = m.group(1)
                        logg.append(f"{namn} {c.coordinate}: bild-URL gick inte att hämta")
            d = ws.cell(row=c.row + offset, column=c.column, value=v)
            if c.has_style:
                d.font = copy(c.font)
                d.fill = copy(c.fill)
                d.border = copy(c.border)
                d.alignment = copy(c.alignment)
                d.number_format = "0.00" if rattad else c.number_format
            if kommentar:
                d.comment = Comment(kommentar, "Master sheet")
                d.fill = PatternFill("solid", fgColor="FFF4B183")  # orange = osäker
                osakra += 1
            if c.hyperlink is not None and c.hyperlink.target:
                d.hyperlink = c.hyperlink.target
            elif isinstance(v, str) and v.startswith("http"):
                d.hyperlink = v

    for rng in src.merged_cells.ranges:
        ws.merge_cells(start_row=rng.min_row + offset, end_row=rng.max_row + offset,
                       start_column=rng.min_col, end_column=rng.max_col)

    for key, dim in src.column_dimensions.items():
        lo = dim.min or column_index_from_string(key)
        hi = dim.max or lo
        for ci in range(lo, hi + 1):
            if dim.width:
                bredd[ci] = max(bredd.get(ci, 0), dim.width)
    for r, dim in src.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[r + offset].height = dim.height

    for img in src._images:
        xi = XlImage(krymp(PILImage.open(img.ref)))
        a = deepcopy(img.anchor)
        a._from.row += offset
        xi.anchor = a
        ws.add_image(xi)
        n_img += 1

    if namn == "Batch 1":
        nytt = lagg_till_marknader(ws, src, offset, sista, 1, None)
    elif namn == "Batch 2":
        nytt = None  # har inga marknadsblock alls
    else:
        nytt = lagg_till_marknader(ws, src, offset, sista, 3, 4)
    rad = offset + sista + 2  # en tom rad mellan batcherna
    if osakra:
        logg.append(f"{namn}: {osakra} celler i US-blocket är osäkra efter tillbakaräkningen (orange + kommentar)")
    print(f"{namn:15s} rader {offset + 1}–{offset + sista}  bilder={n_img}  nya marknader från kol {get_column_letter(nytt) if nytt else '-'}")

for ci, w in bredd.items():
    ws.column_dimensions[get_column_letter(ci)].width = min(w, 60)

# Innehållsförteckning överst? Nej — Axel vill bara scrolla. Men frys inget och dölj inget.
master.save(UT)
print(f"tidsvärden rättade till decimaltal: {tidsfel}")
print(f"\n{UT}: {rad - 2} rader, {len(ws._images)} bilder, {len(ws.merged_cells.ranges)} sammanslagna")
for l in logg:
    print("  !", l)
