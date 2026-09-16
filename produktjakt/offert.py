#!/usr/bin/env python3
"""Skriver körningens fynd till leverantörens offertark (Axels mall).

    python3 offert.py --fynd korningar/<datum>/fynd.json [--ut korningar/<datum>/offert.xlsx]
    python3 offert.py --fynd korningar/<d1>/fynd.json korningar/<d2>/fynd.json --ut <fil>   # flera rundor i ett ark

Bilderna bäddas in i arket (kolumn A:C) ur heron i korningar/<datum>/v3/bilder/ — Axels krav 2026-09-16:
"det MÅSTE vara bilder med i sheetet". =IMAGE()-formeln visas inte i Excel/Numbers, så den används bara
som reserv när ingen lokal bild finns.

Mallen `mall/offertmall.xlsx` är Axels eget ark, rensat på produkter. Varje produkt blir ett block
om fyra rader (tre kvantitetsnivåer + en tunn skiljerad):

  A  =IMAGE("<bild-url>")   A:C sammanslagna över blockets tre rader
  D  produktnamn på svenska
  E  anteckning till leverantören
  I  SWEDEN
  J  K  L  produktkostnad / frakt / totalt — LÄMNAS TOMMA, det är leverantören som fyller dem
  N  produktlänken (N:P)
  R  kvantitetstrappan 100 / 200 / 300

Prisfälten fylls aldrig av oss: arket är en förfrågan, inte ett facit. Vårt eget räknade inköpspris
står i stället som anteckning, så leverantören ser vad vi utgår från.
"""
import argparse
import copy
import datetime
import json
import os
import re
import shutil

import openpyxl
from openpyxl.drawing.image import Image as XlBild

try:
    from PIL import Image as PilBild
except ImportError:  # utan Pillow faller vi tillbaka på =IMAGE()
    PilBild = None

BILD_MAX_B, BILD_MAX_H = 235, 168   # px — ryms i A:C (18+4+13 tecken) × tre rader à 45 pt
BILD_RADHOJD = 45.0

HERE = os.path.dirname(os.path.abspath(__file__))
MALL = os.path.join(HERE, "mall", "offertmall.xlsx")
FORSTA_RAD = 5
RADER_PER_BLOCK = 4
KOL_SLUT = 47


def svenskt_namn(titel, grupp):
    """Kort svenskt namn ur leverantörens engelska titel. Kan inte bli perfekt — därför står
    originaltiteln kvar i anteckningen, så att leverantören alltid ser vad vi menar."""
    t = re.sub(r"^\s*(\d+\s*(pcs?|pack|st)\b[,\s]*)", "", titel, flags=re.I)
    t = re.sub(r"\s*[-–,|]\s*(free shipping|hot sale|new|2026|dropship).*$", "", t, flags=re.I)
    t = re.sub(r"\s+", " ", t).strip()
    return (t[:70] + "…") if len(t) > 70 else t


def hamta_mallblock(ws):
    block = {"hojder": [], "stilar": []}
    for j in range(RADER_PER_BLOCK):
        rr = FORSTA_RAD + j
        block["hojder"].append(ws.row_dimensions[rr].height)
        block["stilar"].append({c: (copy.copy(ws.cell(row=rr, column=c).font),
                                    copy.copy(ws.cell(row=rr, column=c).fill),
                                    copy.copy(ws.cell(row=rr, column=c).border),
                                    copy.copy(ws.cell(row=rr, column=c).alignment),
                                    ws.cell(row=rr, column=c).number_format)
                                for c in range(1, KOL_SLUT + 1)})
    return block


def skriv_block(ws, i, p, mallblock):
    r = FORSTA_RAD + i * RADER_PER_BLOCK
    topp, bott = r, r + 2
    for a, b in (("A", "C"), ("E", "G"), ("H", "H"), ("I", "I"), ("M", "M"), ("N", "P"), ("Q", "Q")):
        ws.merge_cells(f"{a}{topp}:{b}{bott}")
    inbaddad = False
    if p.get("bild_fil") and PilBild and os.path.exists(p["bild_fil"]):
        try:
            im = PilBild.open(p["bild_fil"]).convert("RGB")
            im.thumbnail((BILD_MAX_B, BILD_MAX_H))
            tumme = os.path.splitext(p["bild_fil"])[0] + "-ark.png"
            im.save(tumme, "PNG")
            xb = XlBild(tumme)
            xb.anchor = f"A{topp}"
            ws.add_image(xb)
            inbaddad = True
        except Exception as e:  # trasig bildfil → formeln som reserv
            print(f"  bild kunde inte bäddas in för {p['namn_sv'][:40]}: {e}")
    if not inbaddad and p.get("bild"):
        ws.cell(row=topp, column=1).value = f'=IMAGE("{p["bild"]}")'
    ws.cell(row=topp, column=4).value = p["namn_sv"]
    ws.cell(row=topp, column=5).value = p["notering"]
    ws.cell(row=topp, column=9).value = "SWEDEN"
    ws.cell(row=topp, column=14).value = p["url"]
    for j, qty in enumerate((100, 200, 300)):
        ws.cell(row=topp + j, column=18).value = qty
    for j in range(RADER_PER_BLOCK):
        ws.row_dimensions[r + j].height = (BILD_RADHOJD if (inbaddad and j < 3) else mallblock["hojder"][j])
        for c in range(1, KOL_SLUT + 1):
            st = mallblock["stilar"][j].get(c)
            if st:
                cell = ws.cell(row=r + j, column=c)
                cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = st


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fynd", required=True, nargs="+")
    ap.add_argument("--ut")
    ap.add_argument("--mall", default=MALL)
    a = ap.parse_args()

    filer = a.fynd
    d = json.load(open(filer[-1], encoding="utf-8"))
    datum = d.get("datum") or datetime.date.today().isoformat()
    ut = a.ut or os.path.join(os.path.dirname(filer[-1]), f"Leverantorsoffert-{datum}.xlsx")

    produkter = []
    alla = []
    for f in filer:
        dd = json.load(open(f, encoding="utf-8"))
        for p in dd["produkter"]:
            p["_datum"] = dd.get("datum", "")
            alla.append(p)
    for p in alla:
        e = p["ekonomi"]
        bild_fil = p.get("bild_fil")
        if bild_fil and not os.path.isabs(bild_fil):
            bild_fil = os.path.join(HERE, bild_fil)
        produkter.append({
            "namn_sv": (p.get("namn_sv") or svenskt_namn(p["titel"], p.get("grupp", ""))) + (f"  ({p['_datum'][5:]})" if len(filer) > 1 else ""),
            "url": p["url"], "bild": p.get("bild"), "bild_fil": bild_fil,
            "notering": (f"AliExpress: {p['pris_text']}. Vår räkning: landad ca {e['landad']} kr, "
                         f"tänkt butikspris {e['forslag_pris']} kr. Originaltitel: {p['titel'][:150]}"),
        })

    shutil.copy(a.mall, ut)
    wb = openpyxl.load_workbook(ut)
    ws = wb["Supplier Quote"]
    mallblock = hamta_mallblock(ws)
    for i, p in enumerate(produkter):
        skriv_block(ws, i, p, mallblock)
    slut = FORSTA_RAD + len(produkter) * RADER_PER_BLOCK
    if ws.max_row >= slut:
        ws.delete_rows(slut, ws.max_row - slut + 1)
    ws["A1"] = len(produkter)
    wb.save(ut)

    print(f"{ut}: {len(produkter)} produkter, {os.path.getsize(ut) // 1024} KB")
    for i, p in enumerate(produkter):
        print(f"  rad {FORSTA_RAD + i * RADER_PER_BLOCK:>3}  {p['namn_sv'][:56]}")
    return ut


if __name__ == "__main__":
    main()
