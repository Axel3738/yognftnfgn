# Bygger cogs-matstrumpor.xlsx ur leverantörens offert (skärmdumpar från Axel 2026-09-29).
# Upplägg som Axels quote-sheet: produkterna nedåt, marknaderna bredvid varandra åt höger.
# Kör: python3 matstrumpor/cogs-offert/bygg.py  (kräver openpyxl + pillow)
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties

HAR = os.path.dirname(os.path.abspath(__file__))
BILD = HAR + '/bild/'
UT = HAR + '/cogs-matstrumpor.xlsx'
DATUM = '2026-09-29'

F = 'Arial'
NAVY = '1F3B57'
KREM = PatternFill('solid', fgColor='FFF2CC')
GUL = PatternFill('solid', fgColor='FFF7DB')
TOTAL = PatternFill('solid', fgColor='EEF2F7')
VIT = PatternFill('solid', fgColor='FFFFFF')
LINJE = Side(style='thin', color='D9D9D9')
TJOCK = Side(style='medium', color='94A3B8')
MITT = Alignment(horizontal='center', vertical='center', wrap_text=True)
TAL = '0.00;-0.00;""'
BLA_TEXT = '0000FF'   # blå = siffra ur leverantörens offert

# Marknaderna i ordning, med bandfärg (som exemplets NORWAY/DENMARK/UK/US/CANADA-band)
LANDER = [('SWEDEN', '2F5597'), ('NORWAY', '5B7FB5'), ('DENMARK', 'C9892F'), ('FINLAND', '3E8E7E'),
          ('US', '8E7CC3'), ('UK', 'B45F5F'), ('CANADA', '6FA8DC'), ('AUSTRALIA', 'B8860B'),
          ('NEW ZEALAND', '4F7942')]
TOMMA_LANDER = 8   # förberedda block åt höger för nya marknader
NY_FARG = 'A6A6A6'

# Offerten: produkt -> rader (variant, antal set, {land: (kostnad, frakt, total)})
def rad(variant, qty, **lander):
    return (variant, qty, lander)
PRODUKTER = [
    ('Sushi socks', 'sushi.png', [
        rad('3 pairs', 1, SWEDEN=(1.7, 5.2, 6.9), NORWAY=(1.7, 6.3, 8.0), DENMARK=(2.6, 5.8, 8.4), FINLAND=(2.6, 6.3, 8.9)),
        rad('3 pairs', 2, SWEDEN=(3.4, 7.0, 10.4), NORWAY=(3.4, 7.9, 11.3), DENMARK=(5.2, 6.2, 11.4), FINLAND=(5.2, 7.5, 12.7)),
        rad('5 pairs', 1, SWEDEN=(2.6, 5.6, 8.2), NORWAY=(2.6, 6.8, 9.4), DENMARK=(2.6, 7.1, 9.7), FINLAND=(2.6, 7.7, 10.3)),
        rad('5 pairs', 2, SWEDEN=(5.2, 7.1, 12.3), NORWAY=(5.2, 8.5, 13.7), DENMARK=(5.2, 8.4, 13.6), FINLAND=(5.2, 10.1, 15.3)),
    ]),
    ('Donut socks', 'donut.png', [
        rad('3 pairs', 1, SWEDEN=(2.1, 5.2, 7.3), US=(2.6, 6.5, 9.1), UK=(2.6, 4.5, 7.1), **{'NEW ZEALAND': (2.6, 6.2, 8.7)},
            AUSTRALIA=(2.6, 6.5, 9.1), CANADA=(2.6, 5.5, 8.1), NORWAY=(2.6, 6.6, 9.2), DENMARK=(2.6, 6.8, 9.4), FINLAND=(2.6, 7.3, 9.9)),
        rad('3 pairs', 2, SWEDEN=(4.2, 7.0, 11.3), US=(5.2, 9.0, 14.2), UK=(5.2, 6.2, 11.4), **{'NEW ZEALAND': (5.2, 9.4, 14.5)},
            AUSTRALIA=(5.2, 9.8, 14.9), CANADA=(5.2, 7.5, 12.6), NORWAY=(5.2, 8.8, 13.9), DENMARK=(5.2, 8.9, 14.0), FINLAND=(5.2, 9.4, 14.5)),
    ]),
    ('Pizza socks', 'pizza.png', [
        rad('4 pairs', 1, SWEDEN=(2.7, 5.4, 8.1), US=(2.7, 7.1, 9.8), UK=(2.7, 4.9, 7.6), **{'NEW ZEALAND': (2.7, 7.2, 9.9)},
            AUSTRALIA=(2.7, 7.6, 10.3), CANADA=(2.7, 6.1, 8.8), NORWAY=(2.7, 6.8, 9.5), DENMARK=(2.7, 6.9, 9.6), FINLAND=(2.7, 7.4, 10.1)),
        rad('4 pairs', 2, SWEDEN=(5.4, 7.7, 13.1), US=(5.4, 10.7, 16.1), UK=(5.4, 7.0, 12.4), **{'NEW ZEALAND': (5.4, 11.4, 16.8)},
            AUSTRALIA=(5.4, 11.8, 17.2), CANADA=(5.4, 8.9, 14.3), NORWAY=(5.4, 8.9, 14.3), DENMARK=(5.4, 8.9, 14.3), FINLAND=(5.4, 10.5, 15.9)),
    ]),
    ('Hamburger socks', 'burger.png', [
        rad('3 pairs', 1, SWEDEN=(2.5, 5.2, 7.6), US=(2.5, 6.4, 8.9), UK=(2.5, 4.6, 7.0), **{'NEW ZEALAND': (2.5, 6.4, 8.9)},
            AUSTRALIA=(2.5, 6.8, 9.2), CANADA=(2.5, 5.7, 8.2), NORWAY=(2.5, 6.4, 8.9), DENMARK=(2.5, 6.5, 9.0), FINLAND=(2.5, 7.1, 9.5)),
        rad('3 pairs', 2, SWEDEN=(4.9, 7.3, 12.2), US=(4.9, 9.5, 14.4), UK=(4.9, 6.4, 11.3), **{'NEW ZEALAND': (4.9, 9.9, 14.8)},
            AUSTRALIA=(4.9, 10.3, 15.2), CANADA=(4.9, 7.9, 12.8), NORWAY=(4.9, 8.3, 13.2), DENMARK=(4.9, 8.4, 13.3), FINLAND=(4.9, 9.8, 14.7)),
    ]),
]
TOMMA_PRODUKTER = 3   # förberedda block nedåt för nya produkter (2 rader var)

wb = Workbook()
ws = wb.active
ws.title = 'COGS per market'
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 90

def c(r, k, v=None, font=None, fill=None, fmt=None, align=MITT, border=True):
    x = ws.cell(r, k, v)
    x.font = font or Font(name=F, size=10)
    if fill: x.fill = fill
    if fmt: x.number_format = fmt
    x.alignment = align
    if border: x.border = Border(left=LINJE, right=LINJE, top=LINJE, bottom=LINJE)
    return x

# Kolumner: A marginal | B bild | C produkt | D variant | E antal set | sedan 3 kolumner per marknad
FAST = ['Product image', 'Product name', 'Variant', 'Qty (sets)']
FORSTA = 6
BLOCK = 3
antal_block = len(LANDER) + TOMMA_LANDER
SISTA = FORSTA + antal_block * BLOCK - 1

# Rad 1: mörk rubrikrad (som "BATCH 6")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=SISTA)
c(1, 1, '▶  MATSTRUMPOR — COGS PER MARKET (USD per order)', Font(name=F, size=13, bold=True, color='FFFFFF'),
  PatternFill('solid', fgColor=NAVY), align=Alignment(horizontal='left', vertical='center', indent=1), border=False)
ws.row_dimensions[1].height = 30
# Rad 2: instruktion
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=SISTA)
c(2, 1, f'Supplier quote read {DATUM}. Blue numbers = supplier quote. New market: type the country in a grey header to the right, '
        'then fill product cost and shipping cost (yellow). Total fills itself — overwrite it if the supplier quotes a different total.',
  Font(name=F, size=9, italic=True, color='6B7280'), align=Alignment(horizontal='center', vertical='center', wrap_text=True), border=False)
ws.row_dimensions[2].height = 28

# Rad 3: marknadsband, rad 4: kolumnrubriker
HR1, HR2 = 3, 4
for i, h in enumerate(FAST):
    k = 2 + i
    ws.merge_cells(start_row=HR1, start_column=k, end_row=HR2, end_column=k)
    c(HR1, k, h, Font(name=F, size=10, bold=True, color='FFFFFF'), PatternFill('solid', fgColor=NAVY))
    c(HR2, k, None, fill=PatternFill('solid', fgColor=NAVY))
land_kol = {}
for b in range(antal_block):
    k = FORSTA + b * BLOCK
    namn, farg = LANDER[b] if b < len(LANDER) else (None, NY_FARG)
    ws.merge_cells(start_row=HR1, start_column=k, end_row=HR1, end_column=k + BLOCK - 1)
    c(HR1, k, namn if namn else 'NEW MARKET (type here)',
      Font(name=F, size=11 if namn else 9, bold=bool(namn), italic=not namn, color='FFFFFF'), PatternFill('solid', fgColor=farg))
    for j, h in enumerate(['Product cost', 'Shipping cost', 'Total ex. tax']):
        c(HR2, k + j, h, Font(name=F, size=9, bold=True, color='374151'),
          PatternFill('solid', fgColor='DCE3EC') if j == 2 else KREM)
    if namn: land_kol[namn] = k
ws.row_dimensions[HR1].height = 26
ws.row_dimensions[HR2].height = 30

def blockgrans(r):
    for k in range(2, SISTA + 1):
        x = ws.cell(r, k)
        x.border = Border(left=LINJE, right=LINJE, top=LINJE, bottom=TJOCK)

r = HR2 + 1
avvik = 0
for pi, (namn, bild, rader) in enumerate(PRODUKTER + [(None, None, [('', 1, {}), ('', 2, {})])] * TOMMA_PRODUKTER):
    topp = r
    hojd = 30 if len(rader) > 2 else 56
    for variant, qty, data in rader:
        ny = namn is None
        c(r, 4, variant or None, Font(name=F, size=10), GUL if ny else None)
        c(r, 5, qty, Font(name=F, size=11, bold=True, color='374151'), PatternFill('solid', fgColor='E8ECF4'))
        for b in range(antal_block):
            k = FORSTA + b * BLOCK
            land = LANDER[b][0] if b < len(LANDER) else None
            v = data.get(land) if land else None
            if v:
                kost, frakt, tot = v
                c(r, k, kost, Font(name=F, size=10, color=BLA_TEXT), fmt=TAL)
                c(r, k + 1, frakt, Font(name=F, size=10, color=BLA_TEXT), fmt=TAL)
                x = c(r, k + 2, tot, Font(name=F, size=10, bold=True, color=BLA_TEXT), TOTAL, fmt=TAL)
                if abs(kost + frakt - tot) > 0.051:
                    avvik += 1
                    x.comment = Comment(f'Supplier total as quoted. Cost + shipping = {kost + frakt:.1f} (supplier rounding).', 'COGS')
            else:
                a, s = L(k), L(k + 1)
                c(r, k, None, Font(name=F, size=10, color=BLA_TEXT), GUL, fmt=TAL)
                c(r, k + 1, None, Font(name=F, size=10, color=BLA_TEXT), GUL, fmt=TAL)
                c(r, k + 2, f'=IF(COUNT({a}{r}:{s}{r})=2,{a}{r}+{s}{r},"")', Font(name=F, size=10, bold=True), TOTAL, fmt=TAL)
        ws.row_dimensions[r].height = hojd
        r += 1
    for k in (2, 3):
        ws.merge_cells(start_row=topp, start_column=k, end_row=r - 1, end_column=k)
    c(topp, 2, None, fill=GUL if namn is None else VIT)
    c(topp, 3, namn, Font(name=F, size=11, bold=True, color='111827'), GUL if namn is None else None)
    for rr in range(topp, r):
        for k in (2, 3): ws.cell(rr, k).border = Border(left=LINJE, right=LINJE, top=LINJE, bottom=LINJE)
    blockgrans(r - 1)
    if bild:
        img = XLImage(BILD + bild)
        hojd_px = (r - topp) * hojd * 96 / 72 - 8
        sk = min(150 / img.width, hojd_px / img.height)
        img.width, img.height = int(img.width * sk), int(img.height * sk)
        ws.add_image(img, f'B{topp}')

# Tunn vertikal linje mellan marknaderna
for b in range(antal_block):
    k = FORSTA + b * BLOCK
    for rr in range(HR2, r):
        x = ws.cell(rr, k)
        x.border = Border(left=TJOCK, right=x.border.right, top=x.border.top, bottom=x.border.bottom)

ws.column_dimensions['A'].width = 2
ws.column_dimensions['B'].width = 22
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 11
ws.column_dimensions['E'].width = 9
for b in range(antal_block):
    k = FORSTA + b * BLOCK
    ws.column_dimensions[L(k)].width = 10
    ws.column_dimensions[L(k + 1)].width = 10
    ws.column_dimensions[L(k + 2)].width = 11
ws.freeze_panes = ws.cell(HR2 + 1, FORSTA)
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f'{HR1}:{HR2}'

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(UT)
print(UT, 'avvikande totaler:', avvik, 'sista rad:', r - 1, 'sista kolumn:', L(SISTA))
