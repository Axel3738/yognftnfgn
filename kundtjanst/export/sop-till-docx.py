#!/usr/bin/env python3
"""Markdown → .docx för SOP-mappen. Täcker den delmängd filerna använder:
rubriker, tabeller, punkt- och nummerlistor, checkboxar, citat, kodblock,
fetstil/kursiv/kod/genomstruket, länkar, horisontella linjer.

HTML-kommentarer (granskningsnoterna i slutet av 00-MASTER) tas bort — de är
arbetsanteckningar och ska inte in i ett dokument någon läser.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

GRON = RGBColor(0x16, 0x65, 0x34)
GRA = RGBColor(0x44, 0x44, 0x44)
CHECK = {'[x]': '☑', '[X]': '☑', '[ ]': '☐'}

# ---------------------------------------------------------------- inline

INLINE = re.compile(
    r'(`[^`]+`)'                      # kod
    r'|(\*\*\*[^*]+\*\*\*)'           # fet kursiv
    r'|(\*\*(?:(?!\*\*).)+?\*\*)'     # fet — får bära kursiv inuti sig
    r'|(~~[^~]+~~)'                   # genomstruket
    r'|(\*[^*\n]+\*)'                 # kursiv
    r'|(\[[^\]]+\]\([^)]+\))'         # länk
)


def skriv_inline(p, text, fet=False, kursiv=False):
    """Lägger texten i paragrafen p som runs med rätt formatering."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _run(p, text[pos:m.start()], fet, kursiv)
        bit = m.group(0)
        if bit.startswith('`'):
            r = _run(p, bit[1:-1], fet, kursiv)
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
        elif bit.startswith('***'):
            _run(p, bit[3:-3], True, True)
        elif bit.startswith('**'):
            skriv_inline(p, bit[2:-2], True, kursiv)   # kursiv/kod inuti fetstil
        elif bit.startswith('~~'):
            r = _run(p, bit[2:-2], fet, kursiv)
            r.font.strike = True
        elif bit.startswith('*'):
            _run(p, bit[1:-1], fet, True)
        else:  # [text](url)
            txt, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', bit).groups()
            r = _run(p, txt, fet, kursiv)
            r.font.underline = True
            r.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
            if url.startswith('http') and txt not in url:
                _run(p, f' ({url})', fet, kursiv).font.size = Pt(9)
        pos = m.end()
    if pos < len(text):
        _run(p, text[pos:], fet, kursiv)


def _run(p, txt, fet, kursiv):
    r = p.add_run(txt)
    r.bold = fet
    r.italic = kursiv
    return r


def checkbox(txt):
    return re.sub(r'^\[([ xX])\]\s*', lambda m: CHECK[m.group(0).strip()] + ' ', txt.strip())


def skugga(cell, farg='E8F0E8'):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), farg)
    cell._tc.get_or_add_tcPr().append(sh)


# ---------------------------------------------------------------- block

def lagg_till(doc, md, rubriknivå=0):
    """Skriver markdown i doc. rubriknivå=1 skjuter alla rubriker ett steg ner."""
    md = re.sub(r'<!--.*?-->', '', md, flags=re.S).replace('\r\n', '\n')
    rader = md.split('\n')
    i = 0
    while i < len(rader):
        rad = rader[i]

        # kodblock
        if rad.lstrip().startswith('```'):
            i += 1
            block = []
            while i < len(rader) and not rader[i].lstrip().startswith('```'):
                block.append(rader[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run('\n'.join(block))
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            continue

        # tabell
        if rad.lstrip().startswith('|') and i + 1 < len(rader) \
           and re.fullmatch(r'\s*\|[\s:|-]+\|\s*', rader[i + 1] or ''):
            delar = lambda r: [c for c in r.strip().strip('|').split('|')]
            huvud = delar(rad)
            i += 2
            kropp = []
            while i < len(rader) and rader[i].lstrip().startswith('|'):
                rd = delar(rader[i])
                rd = (rd + [''] * len(huvud))[:len(huvud)]
                kropp.append(rd)
                i += 1
            t = doc.add_table(rows=1, cols=len(huvud))
            t.style = 'Table Grid'
            t.autofit = True
            for c, txt in zip(t.rows[0].cells, huvud):
                c.text = ''
                skugga(c)
                skriv_inline(c.paragraphs[0], checkbox(txt), fet=True)
            for rd in kropp:
                celler = t.add_row().cells
                for c, txt in zip(celler, rd):
                    c.text = ''
                    skriv_inline(c.paragraphs[0], checkbox(txt))
            for r_ in t.rows:
                for c in r_.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        for run in p.runs:
                            if run.font.size is None:
                                run.font.size = Pt(9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # rubrik
        m = re.match(r'(#{1,6})\s+(.*)$', rad)
        if m:
            niva = min(len(m.group(1)) + rubriknivå, 9)
            p = doc.add_heading('', level=niva)
            skriv_inline(p, m.group(2))
            for run in p.runs:
                run.font.color.rgb = GRON
            i += 1
            continue

        # horisontell linje
        if re.fullmatch(r'\s*(-{3,}|\*{3,}|_{3,})\s*', rad):
            p = doc.add_paragraph()
            pb = OxmlElement('w:pBdr')
            b = OxmlElement('w:bottom')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '6')
            b.set(qn('w:color'), 'CCCCCC')
            pb.append(b)
            p._p.get_or_add_pPr().append(pb)
            i += 1
            continue

        # citat
        if rad.lstrip().startswith('>'):
            block = []
            while i < len(rader) and rader[i].lstrip().startswith('>'):
                block.append(re.sub(r'^\s*>\s?', '', rader[i]))
                i += 1
            # Radbrytning inuti ett citat är ingen ny paragraf — utan den här
            # hopslagningen bryts **fetstil** som spänner över två rader, och
            # stjärnorna hamnar i texten.
            stycken, aktuellt = [], []
            for b in block:
                if not b.strip():
                    if aktuellt:
                        stycken.append(' '.join(aktuellt))
                        aktuellt = []
                elif re.match(r'\s*([-*•]|\d+[.)])\s+', b) or b.lstrip().startswith('|'):
                    if aktuellt:
                        stycken.append(' '.join(aktuellt))
                        aktuellt = []
                    stycken.append(b.strip())
                else:
                    aktuellt.append(b.strip())
            if aktuellt:
                stycken.append(' '.join(aktuellt))

            for b in stycken:
                if not b.strip():
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(3)
                pb = OxmlElement('w:pBdr')
                left = OxmlElement('w:left')
                left.set(qn('w:val'), 'single')
                left.set(qn('w:sz'), '18')
                left.set(qn('w:space'), '8')
                left.set(qn('w:color'), '166534')
                pb.append(left)
                p._p.get_or_add_pPr().append(pb)
                skriv_inline(p, checkbox(re.sub(r'^[-*]\s+', '• ', b.strip())))
            continue

        # lista
        m = re.match(r'(\s*)([-*•]|\d+[.)])\s+(.*)$', rad)
        if m:
            djup = len(m.group(1).replace('\t', '  ')) // 2
            numrerad = m.group(2)[0].isdigit()
            stil = 'List Number' if numrerad else 'List Bullet'
            if djup:
                stil = f'{stil} {min(djup + 1, 3)}'
            try:
                p = doc.add_paragraph(style=stil)
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25 * (djup + 1))
            p.paragraph_format.space_after = Pt(3)
            # Ombruten listrad hör till samma punkt — annars hamnar fortsättningen
            # som egen paragraf och **fetstil** över två rader går sönder.
            txt = [m.group(3)]
            i += 1
            while i < len(rader) and rader[i].strip() and not re.match(
                    r'\s*(#{1,6}\s|[-*•]\s|\d+[.)]\s|>|\||```|-{3,}|_{3,})', rader[i]):
                txt.append(rader[i].strip())
                i += 1
            skriv_inline(p, checkbox(' '.join(txt)))
            continue

        # tom rad
        if not rad.strip():
            i += 1
            continue

        # paragraf (flera rader i följd)
        block = [rad]
        i += 1
        while i < len(rader) and rader[i].strip() and not re.match(
                r'\s*(#{1,6}\s|[-*•]\s|\d+[.)]\s|>|\||```|-{3,}|_{3,})', rader[i]):
            block.append(rader[i])
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        skriv_inline(p, ' '.join(x.strip() for x in block))


def nytt_dokument():
    doc = Document()
    n = doc.styles['Normal']
    n.font.name = 'Calibri'
    n.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.7)
    return doc


ORDNING = [
    'README.md', 'START-HERE.md', '00-MASTER.md',
    '10-NOT-RECEIVED.md', '11-UNACCEPTABLE.md', '12-CREDIT-NOT-PROCESSED.md',
    '13-FRAUD-UNRECOGNIZED.md', '14-DUPLICATE-SUBSCRIPTION-OTHER.md',
    '20-NO-CONTACT.md', '30-EMAIL-TEMPLATES.md', '40-EVIDENCE-PACK.md',
    '50-PREVENTION.md', '60-ESCALATION.md', '99-BACKLOG.md',
]


def main():
    sop = Path(sys.argv[1])
    ut = Path(sys.argv[2])
    (ut / 'enskilda').mkdir(parents=True, exist_ok=True)

    funna = sorted(p.name for p in sop.glob('*.md'))
    saknade = [f for f in funna if f not in ORDNING]
    if saknade:
        sys.exit(f'Filer utanför läsordningen: {saknade}')
    for f in ORDNING:
        if f not in funna:
            sys.exit(f'{f} saknas i {sop}')
    beslut = sorted(sop.glob('beslut/*.md'))

    # En fil per SOP
    for f in ORDNING:
        doc = nytt_dokument()
        lagg_till(doc, (sop / f).read_text(encoding='utf-8'))
        doc.save(ut / 'enskilda' / f.replace('.md', '.docx'))

    # Den samlade handboken
    doc = nytt_dokument()
    t = doc.add_heading('Chargeback SOPs', level=0)
    for run in t.runs:
        run.font.color.rgb = GRON
    p = doc.add_paragraph()
    skriv_inline(p, 'Bäverbutiken och alla butiker på samma backend. '
                    'Byggd ur 50 verkliga tvister 2026-09-20. Procedurtexten är identisk '
                    'på alla butiker — bara `{{PLATSHÅLLARNA}}` byts, ur butikens brandfil.')
    p = doc.add_paragraph()
    r = p.add_run('Källa: kundtjanst/sop/ i repot. Ändrar du något i det här dokumentet '
                  'kommer det inte tillbaka till repot av sig själv — säg vad du ändrat, '
                  'eller be om en ny export efteråt.')
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GRA

    h = doc.add_heading('Innehåll', level=1)
    for run in h.runs:
        run.font.color.rgb = GRON
    for nr, f in enumerate(ORDNING, 1):
        doc.add_paragraph(f'{nr}. {f[:-3]}').paragraph_format.space_after = Pt(2)
    doc.add_paragraph(f'{len(ORDNING) + 1}. Bilaga: beslutsblad per order '
                      f'({len(beslut)} st)').paragraph_format.space_after = Pt(2)

    for f in ORDNING:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        k = doc.add_heading(f[:-3], level=1)
        for run in k.runs:
            run.font.color.rgb = GRON
        lagg_till(doc, (sop / f).read_text(encoding='utf-8'), rubriknivå=1)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    k = doc.add_heading('Bilaga — beslutsblad per order', level=1)
    for run in k.runs:
        run.font.color.rgb = GRON
    doc.add_paragraph('En färdig dom per öppen tvist, med bevistext att klistra in i Shopify.')
    for b in beslut:
        lagg_till(doc, b.read_text(encoding='utf-8'), rubriknivå=1)

    doc.save(ut / 'Chargeback-SOP-handboken.docx')
    print(f'{len(ORDNING)} SOP-filer + {len(beslut)} beslutsblad → {ut}')


if __name__ == '__main__':
    main()
